"""Displacement calibration models and voltage resolution."""

from __future__ import annotations

import csv
import json
import math
import time
from collections.abc import Iterable
from dataclasses import asdict, dataclass
from itertools import pairwise
from pathlib import Path
from statistics import median

from droplet_lab.config import MAX_AMPLITUDE_VPP, DisplacementConfig, ExperimentConfig
from droplet_lab.devices.base import FunctionGenerator, Oscilloscope
from droplet_lab.state import ExperimentState
from droplet_lab.storage import utc_now_filename_safe, utc_now_iso


@dataclass(frozen=True, slots=True)
class VoltageLimitPoint:
    frequency_hz: float
    max_voltage_v: float


@dataclass(frozen=True, slots=True)
class CalibrationMeasurement:
    set_amplitude_vpp: float
    median_scope_vpp_v: float
    median_displacement_um: float
    samples: list[float]


@dataclass(frozen=True, slots=True)
class CalibrationCurve:
    frequency_hz: float
    safe_max_amplitude_vpp: float
    measurements: list[CalibrationMeasurement]


@dataclass(frozen=True, slots=True)
class CalibrationModel:
    version: int
    created_at_utc: str
    vibrometer_factor_um_per_v: float
    amplifier_gain: float
    curves: list[CalibrationCurve]


@dataclass(frozen=True, slots=True)
class ResolvedDisplacement:
    frequency_hz: float
    target_displacement_um: float
    amplitude_vpp: float


@dataclass(frozen=True, slots=True)
class ValidationResult:
    frequency_hz: float
    target_displacement_um: float
    amplitude_vpp: float
    measured_displacement_um: float | None
    abs_error_um: float | None
    percent_error: float | None
    warning: bool


def calibration_frequencies(cfg: DisplacementConfig) -> list[float]:
    out: list[float] = []
    current = cfg.calibration_start_hz
    # Small tolerance keeps decimal steps from missing the endpoint.
    while current <= cfg.calibration_stop_hz + 1e-9:
        out.append(round(current, 10))
        current += cfg.calibration_step_hz
    return out


def voltage_steps(max_amplitude_vpp: float, count: int) -> list[float]:
    if count <= 0:
        raise ValueError("voltage step count must be positive")
    return [max_amplitude_vpp * i / count for i in range(1, count + 1)]


def load_voltage_limits_xlsx(path: Path) -> list[VoltageLimitPoint]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read .xlsx voltage limit files") from exc

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"{path}: workbook is empty")
    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    try:
        frequency_idx = headers.index("Frequency [Hz]")
        max_voltage_idx = headers.index("max. Voltage [V]")
    except ValueError as exc:
        raise ValueError(
            f"{path}: expected headers 'Frequency [Hz]' and 'max. Voltage [V]'"
        ) from exc

    points: list[VoltageLimitPoint] = []
    for row in rows[1:]:
        if frequency_idx >= len(row) or max_voltage_idx >= len(row):
            continue
        raw_freq = row[frequency_idx]
        raw_max = row[max_voltage_idx]
        if raw_freq is None or raw_max is None:
            continue
        points.append(VoltageLimitPoint(frequency_hz=float(raw_freq), max_voltage_v=float(raw_max)))
    if len(points) < 2:
        raise ValueError(f"{path}: expected at least two voltage limit rows")
    return sorted(points, key=lambda p: p.frequency_hz)


def safe_amplitude_limit_vpp(
    frequency_hz: float,
    *,
    limits: list[VoltageLimitPoint],
    amplifier_gain: float,
) -> float:
    max_voltage_v = _interpolate_by_frequency(
        frequency_hz,
        [(p.frequency_hz, p.max_voltage_v) for p in limits],
        label="voltage limit",
    )
    return min(MAX_AMPLITUDE_VPP, max_voltage_v / amplifier_gain)


def save_calibration_model(model: CalibrationModel, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(asdict(model), indent=2), encoding="utf-8")


def load_calibration_model(path: Path) -> CalibrationModel:
    raw = json.loads(path.read_text(encoding="utf-8"))
    curves = [
        CalibrationCurve(
            frequency_hz=float(curve["frequency_hz"]),
            safe_max_amplitude_vpp=float(curve["safe_max_amplitude_vpp"]),
            measurements=[
                CalibrationMeasurement(
                    set_amplitude_vpp=float(measurement["set_amplitude_vpp"]),
                    median_scope_vpp_v=float(measurement["median_scope_vpp_v"]),
                    median_displacement_um=float(measurement["median_displacement_um"]),
                    samples=[float(v) for v in measurement.get("samples", [])],
                )
                for measurement in curve["measurements"]
            ],
        )
        for curve in raw["curves"]
    ]
    model = CalibrationModel(
        version=int(raw["version"]),
        created_at_utc=str(raw["created_at_utc"]),
        vibrometer_factor_um_per_v=float(raw["vibrometer_factor_um_per_v"]),
        amplifier_gain=float(raw["amplifier_gain"]),
        curves=curves,
    )
    validate_calibration_model(model)
    return model


def validate_calibration_model(model: CalibrationModel) -> None:
    if model.version != 1:
        raise ValueError(f"unsupported displacement calibration model version {model.version}")
    if len(model.curves) < 1:
        raise ValueError("displacement calibration model has no frequency curves")
    seen_freqs: set[float] = set()
    for curve in model.curves:
        if curve.frequency_hz in seen_freqs:
            raise ValueError(f"duplicate calibration frequency {curve.frequency_hz}")
        seen_freqs.add(curve.frequency_hz)
        if len(curve.measurements) < 2:
            raise ValueError(f"frequency {curve.frequency_hz} needs at least two measurements")
        previous_disp = -math.inf
        previous_amp = -math.inf
        for measurement in sorted(curve.measurements, key=lambda m: m.set_amplitude_vpp):
            if measurement.set_amplitude_vpp <= previous_amp:
                raise ValueError(f"frequency {curve.frequency_hz} has duplicate voltage points")
            if measurement.median_displacement_um <= previous_disp:
                raise ValueError(
                    f"frequency {curve.frequency_hz} displacement is not monotonic with voltage"
                )
            previous_amp = measurement.set_amplitude_vpp
            previous_disp = measurement.median_displacement_um


def resolve_displacement(
    model: CalibrationModel,
    *,
    frequency_hz: float,
    target_displacement_um: float,
) -> float:
    validate_calibration_model(model)
    curves = sorted(model.curves, key=lambda c: c.frequency_hz)
    if frequency_hz < curves[0].frequency_hz or frequency_hz > curves[-1].frequency_hz:
        raise ValueError(
            f"frequency {frequency_hz} Hz is outside calibrated range "
            f"{curves[0].frequency_hz}..{curves[-1].frequency_hz} Hz"
        )

    lower = curves[0]
    upper = curves[-1]
    for index, curve in enumerate(curves):
        if curve.frequency_hz == frequency_hz:
            amplitude = _resolve_on_curve(curve, target_displacement_um)
            _ensure_safe_amplitude(amplitude, curve)
            return amplitude
        if curve.frequency_hz > frequency_hz:
            lower = curves[index - 1]
            upper = curve
            break

    lower_amp = _resolve_on_curve(lower, target_displacement_um)
    upper_amp = _resolve_on_curve(upper, target_displacement_um)
    amplitude = _linear_interpolate(
        frequency_hz,
        lower.frequency_hz,
        lower_amp,
        upper.frequency_hz,
        upper_amp,
    )
    interpolated_safe = _linear_interpolate(
        frequency_hz,
        lower.frequency_hz,
        lower.safe_max_amplitude_vpp,
        upper.frequency_hz,
        upper.safe_max_amplitude_vpp,
    )
    if amplitude > interpolated_safe:
        raise ValueError(
            f"resolved amplitude {amplitude:.4g} Vpp exceeds safe limit "
            f"{interpolated_safe:.4g} Vpp at {frequency_hz:g} Hz"
        )
    return amplitude


def resolve_sweep_displacements(cfg: ExperimentConfig) -> dict[tuple[float, float], float] | None:
    if cfg.sweep.displacements_um is None:
        return None
    if cfg.displacement.model_path is None:
        raise ValueError("displacement.model_path is required")
    model = load_calibration_model(cfg.displacement.model_path)
    return {
        (float(freq), float(disp)): resolve_displacement(
            model,
            frequency_hz=float(freq),
            target_displacement_um=float(disp),
        )
        for freq in cfg.sweep.frequencies_hz
        for disp in cfg.sweep.displacements_um
    }


def run_displacement_calibration(
    *,
    cfg: ExperimentConfig,
    fg: FunctionGenerator,
    scope: Oscilloscope,
    state: ExperimentState | None = None,
) -> tuple[CalibrationModel, Path, Path]:
    displacement = cfg.displacement
    if displacement.max_voltage_table_path is None:
        raise ValueError("displacement.max_voltage_table_path is required for calibration")
    limits = load_voltage_limits_xlsx(displacement.max_voltage_table_path)
    curves: list[CalibrationCurve] = []
    fg.set_sine()
    fg.enable_output(False)
    for frequency_hz in calibration_frequencies(displacement):
        safe_limit = safe_amplitude_limit_vpp(
            frequency_hz,
            limits=limits,
            amplifier_gain=displacement.amplifier_gain,
        )
        fg.set_frequency_hz(frequency_hz)
        fg.enable_output(True)
        measurements: list[CalibrationMeasurement] = []
        for amplitude_vpp in voltage_steps(safe_limit, displacement.voltage_steps):
            fg.set_amplitude_vpp(amplitude_vpp)
            if state is not None:
                state.update(
                    combo_index=0,
                    set_speed_rpm=0,
                    set_frequency_hz=frequency_hz,
                    set_amplitude_vpp=amplitude_vpp,
                    target_displacement_um=None,
                )
            samples = _measure_displacement_samples(
                scope=scope,
                vibrometer_factor_um_per_v=cfg.vibrometer.factor_um_per_v,
                measurement_s=displacement.measurement_s,
                interval_s=displacement.scope_interval_s,
            )
            measurements.append(
                CalibrationMeasurement(
                    set_amplitude_vpp=amplitude_vpp,
                    median_scope_vpp_v=median(samples) / cfg.vibrometer.factor_um_per_v,
                    median_displacement_um=median(samples),
                    samples=samples,
                )
            )
        curves.append(
            CalibrationCurve(
                frequency_hz=frequency_hz,
                safe_max_amplitude_vpp=safe_limit,
                measurements=measurements,
            )
        )
    fg.enable_output(False)
    model = CalibrationModel(
        version=1,
        created_at_utc=utc_now_iso(),
        vibrometer_factor_um_per_v=cfg.vibrometer.factor_um_per_v,
        amplifier_gain=displacement.amplifier_gain,
        curves=curves,
    )
    validate_calibration_model(model)
    model_path = _default_model_path(cfg)
    csv_path = model_path.with_suffix(".csv")
    save_calibration_model(model, model_path)
    write_calibration_csv(model, csv_path)
    return model, model_path, csv_path


def validate_displacement_targets(
    *,
    cfg: ExperimentConfig,
    fg: FunctionGenerator,
    scope: Oscilloscope,
    resolved: dict[tuple[float, float], float],
    state: ExperimentState | None = None,
) -> list[ValidationResult]:
    if cfg.sweep.displacements_um is None:
        return []
    fg.set_sine()
    fg.enable_output(False)
    results: list[ValidationResult] = []
    for frequency_hz, target_um in _unique_frequency_displacements(
        cfg.sweep.frequencies_hz, cfg.sweep.displacements_um
    ):
        amplitude_vpp = resolved[(frequency_hz, target_um)]
        fg.set_frequency_hz(frequency_hz)
        fg.set_amplitude_vpp(amplitude_vpp)
        fg.enable_output(True)
        if state is not None:
            state.update(
                combo_index=0,
                set_speed_rpm=0,
                set_frequency_hz=frequency_hz,
                set_amplitude_vpp=amplitude_vpp,
                target_displacement_um=target_um,
            )
        samples = _measure_displacement_samples(
            scope=scope,
            vibrometer_factor_um_per_v=cfg.vibrometer.factor_um_per_v,
            measurement_s=cfg.displacement.measurement_s,
            interval_s=cfg.displacement.scope_interval_s,
        )
        measured_um = median(samples) if samples else None
        abs_error = abs(measured_um - target_um) if measured_um is not None else None
        percent_error = abs_error / target_um * 100.0 if abs_error is not None else None
        warning = (
            percent_error is None or percent_error > cfg.displacement.validation_threshold_percent
        )
        results.append(
            ValidationResult(
                frequency_hz=frequency_hz,
                target_displacement_um=target_um,
                amplitude_vpp=amplitude_vpp,
                measured_displacement_um=measured_um,
                abs_error_um=abs_error,
                percent_error=percent_error,
                warning=warning,
            )
        )
    fg.enable_output(False)
    return results


def write_calibration_csv(model: CalibrationModel, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fp:
        writer = csv.DictWriter(
            fp,
            fieldnames=[
                "frequency_hz",
                "set_amplitude_vpp",
                "median_scope_vpp_v",
                "median_displacement_um",
                "safe_max_amplitude_vpp",
            ],
            delimiter=";",
        )
        writer.writeheader()
        for curve in model.curves:
            for measurement in curve.measurements:
                writer.writerow(
                    {
                        "frequency_hz": curve.frequency_hz,
                        "set_amplitude_vpp": measurement.set_amplitude_vpp,
                        "median_scope_vpp_v": measurement.median_scope_vpp_v,
                        "median_displacement_um": measurement.median_displacement_um,
                        "safe_max_amplitude_vpp": curve.safe_max_amplitude_vpp,
                    }
                )


def _default_model_path(cfg: ExperimentConfig) -> Path:
    if cfg.displacement.model_path is not None:
        return cfg.displacement.model_path
    folder = cfg.output.base_dir / f"{utc_now_filename_safe()}__displacement_calibration"
    return folder / "displacement_calibration.json"


def _measure_displacement_samples(
    *,
    scope: Oscilloscope,
    vibrometer_factor_um_per_v: float,
    measurement_s: float,
    interval_s: float,
) -> list[float]:
    deadline = time.monotonic() + measurement_s
    samples: list[float] = []
    while True:
        measurement = scope.measure()
        if measurement.vpp_v is not None:
            samples.append(measurement.vpp_v * vibrometer_factor_um_per_v)
        if time.monotonic() >= deadline:
            break
        time.sleep(interval_s)
    if not samples:
        raise ValueError("no valid oscilloscope displacement samples captured")
    return samples


def _unique_frequency_displacements(
    frequencies_hz: Iterable[float], displacements_um: Iterable[float]
) -> list[tuple[float, float]]:
    seen: set[tuple[float, float]] = set()
    out: list[tuple[float, float]] = []
    for frequency_hz in frequencies_hz:
        for displacement_um in displacements_um:
            key = (float(frequency_hz), float(displacement_um))
            if key not in seen:
                seen.add(key)
                out.append(key)
    return out


def _resolve_on_curve(curve: CalibrationCurve, target_displacement_um: float) -> float:
    points = sorted(
        (
            (measurement.median_displacement_um, measurement.set_amplitude_vpp)
            for measurement in curve.measurements
        ),
        key=lambda point: point[0],
    )
    return _interpolate_by_frequency(
        target_displacement_um,
        points,
        label=f"displacement at {curve.frequency_hz:g} Hz",
    )


def _ensure_safe_amplitude(amplitude_vpp: float, curve: CalibrationCurve) -> None:
    if amplitude_vpp > curve.safe_max_amplitude_vpp:
        raise ValueError(
            f"resolved amplitude {amplitude_vpp:.4g} Vpp exceeds safe limit "
            f"{curve.safe_max_amplitude_vpp:.4g} Vpp at {curve.frequency_hz:g} Hz"
        )


def _interpolate_by_frequency(
    value: float, points: list[tuple[float, float]], *, label: str
) -> float:
    ordered = sorted(points, key=lambda p: p[0])
    if value < ordered[0][0] or value > ordered[-1][0]:
        raise ValueError(
            f"{label} value {value:g} is outside range {ordered[0][0]:g}..{ordered[-1][0]:g}"
        )
    for x, y in ordered:
        if value == x:
            return y
    for (x0, y0), (x1, y1) in pairwise(ordered):
        if x0 <= value <= x1:
            return _linear_interpolate(value, x0, y0, x1, y1)
    raise ValueError(f"could not interpolate {label} at {value:g}")


def _linear_interpolate(x: float, x0: float, y0: float, x1: float, y1: float) -> float:
    if x1 == x0:
        return y0
    ratio = (x - x0) / (x1 - x0)
    return y0 + ratio * (y1 - y0)
